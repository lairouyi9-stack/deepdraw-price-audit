#!/usr/bin/env python3
"""Build the Deepdraw price-audit workbook from a listing plan and normalized page data."""

from __future__ import annotations

import argparse
import json
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SKU_CHECKS = [
    ("商家SKU", "价格", "PRICE", "tag"), ("商家SKU", "零售价", "RETAIL_PRICE", "tag"),
    ("天猫", "天猫特卖折扣价", "TM_DISCOUNT_PRICE", "tag"), ("天猫", "天猫特卖专柜价", "TM_SHOP_PRICE", "tag"),
    ("京东", "京东价", "JD_PRICE", "tag"), ("京东", "采购价", "PURCHASE_PRICE", "tag"),
    ("京东", "划线价", "LINEATION_PRICE", "tag"), ("京东", "京东自营市场价", "JD_VC_MARKET_PRICE", "tag"),
    ("有赞", "有赞价格", "YOUZAN_PRICE", "tag"), ("有赞", "有赞标准价", "YOUZAN_STANDARD_PRICE", "tag"),
    ("拼多多", "拼多多单买价", "INDIVIDUAL_PRICE", "pdd1"), ("拼多多", "拼多多团购价", "GROUP_PRICE", "pdd2"),
    ("小红书", "小红书市场价", "PRICE_OF_LITTLE_RED_BOOK", "tag"), ("小红书", "原价", "ORIGINAL_PRICE", "tag"),
    ("抖音", "抖音价", "DOUYIN_PRICE", "tag"), ("抖音", "抖音结算价格", "DOUYIN_SETTLEMENT_PRICE", "tag"),
    ("快手", "快手价", "KUAISHOU_NEW_PRICE", "tag"), ("爱库存", "爱库存供货价", "AIKUCUN_SUPPLY_PRICE", "tag"),
    ("爱库存", "爱库存最低价", "AIKUCUN_FLOOR_PRICE", "tag"), ("爱库存", "爱库存零售价", "AIKUCUN_RETAIL_PRICE", "tag"),
    ("好衣库", "好衣库价", "HAOYK_PRICE", "tag"), ("好衣库", "好衣库原价", "HAOYK_ORIGINAL_PRICE", "tag"),
    ("好衣库", "好衣库供货价", "HAOYK_SUPPLY_PRICE", "tag"), ("好衣库", "好衣库结算价", "HAOYK_SETTLEMENT_PRICE", "tag"),
    ("微信视频小店", "微信视频小店价格", "WX_PRICE", "tag"),
]
SINGLE_CHECKS = [
    ("多平台通用", "抖音参考价", "douyinReferencePrice", "tag", "单值"), ("1688", "产品单价", "alibabaUnitPrice", "tag", "价格区间首档"),
    ("天猫", "奥莱店折扣价", "tmallOutletDiscountPrice", "tag", "单值"), ("京东", "京东市场价", "jdMarketPrice", "jd_market", "单值"),
    ("唯品会", "唯品会市场价（页面字段：划线价）", "vipMarketPrice", "blank", "单值"),
    ("拼多多", "商品市场价", "pddMarketPrice", "tag", "单值"), ("天猫", "专柜价", "tmallCounterPrice", "fixed10000", "单值"),
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--plan", required=True, type=Path)
    parser.add_argument("--codes", required=True, type=Path)
    parser.add_argument("--actual", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    return parser.parse_args()


def normalize_code(value: Any) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text[:-2] if text.endswith(".0") else text


def money(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        return Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        return None


def display_number(value: Decimal) -> int | float:
    return int(value) if value == value.to_integral_value() else float(value)


def fmt(value: Decimal) -> str:
    text = format(value.normalize(), "f")
    return text.rstrip("0").rstrip(".") if "." in text else text


def load_codes(path: Path) -> list[str]:
    return list(dict.fromkeys(re.findall(r"d{10,16}", path.read_text(encoding="utf-8"))))


def find_header(row: tuple[Any, ...], names: set[str]) -> int | None:
    for index, value in enumerate(row):
        if str(value or "").replace("
", "").strip() in names:
            return index
    return None


def read_plan(path: Path, codes: list[str]) -> dict[str, dict[str, Any]]:
    wanted = set(codes)
    workbook = load_workbook(path, read_only=True, data_only=True)
    matches: dict[str, list[dict[str, Any]]] = {code: [] for code in codes}
    candidates = [name for name in workbook.sheetnames if re.search(r"期货|线上|全域", name)]
    if not candidates:
        raise ValueError("上市计划表中未找到名称包含‘期货/线上/全域’的工作表")
    for name in candidates:
        sheet = workbook[name]
        indexes = None
        header_row = 0
        for number, row in enumerate(sheet.iter_rows(min_row=1, max_row=20, values_only=True), 1):
            style_col = find_header(row, {"大货款号", "款号", "商品款号"})
            price_col = find_header(row, {"吊牌价"})
            line_col = find_header(row, {"产品线"})
            if None not in (style_col, price_col, line_col):
                indexes, header_row = (style_col, price_col, line_col), number
                break
        if indexes is None:
            continue
        style_col, price_col, line_col = indexes
        for number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
            code = normalize_code(row[style_col] if style_col < len(row) else None)
            if code in wanted:
                matches[code].append({"sheet": name, "row": number,
                    "tag": money(row[price_col] if price_col < len(row) else None),
                    "product_line": str(row[line_col] or "").strip() if line_col < len(row) else ""})
    result, missing, conflicts = {}, [], []
    for code in codes:
        rows = matches[code]
        if not rows:
            missing.append(code)
            continue
        tags = {row["tag"] for row in rows if row["tag"] is not None}
        lines = {row["product_line"] for row in rows if row["product_line"]}
        if len(tags) != 1 or len(lines) != 1:
            conflicts.append(code)
        else:
            result[code] = {"tag": next(iter(tags)), "product_line": next(iter(lines)), "matches": rows}
    if missing:
        raise ValueError("上市计划表未匹配到款号：" + "、".join(missing))
    if conflicts:
        raise ValueError("上市计划表吊牌价或产品线存在冲突：" + "、".join(conflicts))
    return result


def target(kind: str, tag: Decimal, product_line: str) -> Decimal | str:
    if kind == "pdd1": return tag - Decimal("1")
    if kind == "pdd2": return tag - Decimal("2")
    if kind == "fixed10000": return Decimal("10000")
    if kind == "blank" or (kind == "jd_market" and product_line == "鞋品"): return "应为空"
    return tag


def note(ok: bool, kind: str, actual: Decimal | None, expected: Decimal | str, missing: bool, bad_count: int = 0) -> str:
    if ok:
        return {"pdd1": "符合吊牌价-1元规则", "pdd2": "符合吊牌价-2元规则", "blank": "符合留空要求",
                "fixed10000": "符合固定值要求", "jd_market": "符合产品线规则"}.get(kind, "与吊牌价一致")
    if missing: return "页面未展示该字段，无法取得应填价格"
    if actual is None: return "填写值为空"
    if bad_count: return f"{bad_count}个SKU不符合目标值"
    if isinstance(expected, Decimal):
        diff = actual - expected
        return f"{'高于' if diff > 0 else '低于'}目标值{fmt(abs(diff))}元"
    return "应为空但已填写"


def build_rows(codes: list[str], plan: dict[str, dict[str, Any]], actual_data: list[dict[str, Any]]) -> tuple[list[list[Any]], list[int]]:
    actual_map = {str(item.get("code", "")): item for item in actual_data}
    rows, wrong_rows, serial = [], [], 1
    for code in codes:
        tag, product_line = plan[code]["tag"], plan[code]["product_line"]
        record = actual_map.get(code, {})
        singles = record.get("singles", {})
        sku_map = {item.get("code"): item for item in record.get("skuFields", [])}
        def append(platform: str, field: str, value: Any, expected: Any, coverage: str, status: str, explanation: str) -> None:
            nonlocal serial
            rows.append([serial, code, product_line, platform, field, value, expected, coverage, status, explanation])
            if status == "错误": wrong_rows.append(len(rows) + 1)
            serial += 1
        raw = singles.get("productPrice")
        actual = money(raw)
        ok = raw is not None and actual == tag
        append("多平台通用", "商品价格", "未展示" if raw is None else ("空" if raw == "" else display_number(actual)), display_number(tag), "单值", "正确" if ok else "错误", note(ok, "tag", actual, tag, raw is None))
        for platform, field, field_code, kind in SKU_CHECKS:
            info = sku_map.get(field_code)
            expected = target(kind, tag, product_line)
            missing = not info or not info.get("values")
            raw_values = [] if missing else list(info.get("values", []))
            values = [money(value) for value in raw_values]
            bad_count = sum(value != expected for value in values)
            ok = not missing and bool(values) and bad_count == 0
            unique = list(dict.fromkeys(raw_values))
            if missing: value_display, coverage = "未展示", "字段未展示"
            elif len(unique) == 1:
                value_display = "空" if unique[0] == "" else display_number(money(unique[0]))
                coverage = f"{len(raw_values)}个SKU，均为{'空' if unique[0] == '' else fmt(money(unique[0]))}"
            else:
                value_display = "、".join("空" if item == "" else fmt(money(item)) for item in unique)
                coverage = f"{len(raw_values)}个SKU，存在{len(unique)}种填写值"
            append(platform, field, value_display, display_number(expected), coverage, "正确" if ok else "错误", note(ok, kind, values[0] if len(unique) == 1 else None, expected, missing, bad_count))
        for platform, field, key, kind, base_coverage in SINGLE_CHECKS:
            raw = singles.get(key)
            missing = raw is None
            expected = target(kind, tag, product_line)
            actual = money(raw)
            expects_blank = expected == "应为空"
            ok = not missing and ((raw == "") if expects_blank else actual == expected)
            value_display = "未展示" if missing else ("空" if raw == "" else display_number(actual))
            coverage = base_coverage
            if key == "alibabaUnitPrice":
                coverage = f"价格区间首档，购买数量≥{singles.get('alibabaMinQty') or '未填写'}"
            append(platform, field, value_display, expected if expects_blank else display_number(expected), coverage, "正确" if ok else "错误", note(ok, kind, actual, expected, missing))
        tag_price, wash_price = money(record.get("tagImagePrice")), money(record.get("washLabelPrice"))
        tag_coverage = record.get("tagImageCoverage", "吊牌图")
        wash_coverage = record.get("washLabelCoverage", "洗唛图")
        if tag_price is not None:
            image_price, source, coverage = tag_price, "吊牌图", tag_coverage
            status = "正确" if tag_price == tag else "错误"
            explanation = "吊牌图价格与上市计划吊牌价一致" if status == "正确" else f"吊牌图价格{fmt(tag_price)}与上市计划吊牌价{fmt(tag)}不一致"
            if wash_price is not None and wash_price != tag:
                status = "错误"
                explanation += f"；洗唛图价格{fmt(wash_price)}也与计划价冲突"
        elif wash_price is not None:
            image_price, source, coverage = wash_price, "洗唛图", wash_coverage
            status = "正确" if wash_price == tag else "错误"
            explanation = "吊牌图未识别到价格；洗唛图价格与上市计划吊牌价一致，按规则不报错" if status == "正确" else f"吊牌图未识别到价格；洗唛图价格{fmt(wash_price)}与上市计划吊牌价{fmt(tag)}不一致"
        else:
            image_price, source, coverage, status = "未识别到价格", "吊牌图/洗唛图", f"{tag_coverage}；{wash_coverage}", "待确认"
            explanation = "吊牌图和洗唛图均未识别到价格，需人工确认；未据此判定为价格不一致"
        append("深绘素材", f"吊牌/洗唛图价格（识别来源：{source}）", image_price, display_number(tag), coverage, status, explanation)
    return rows, wrong_rows


def write_workbook(rows: list[list[Any]], wrong_rows: list[int], output: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "排查结果"
    sheet.append(["序号", "款号", "产品线", "平台/区域", "检查字段", "填写值", "目标值/要求", "覆盖范围", "核查结果", "说明"])
    for row in rows: sheet.append(row)
    blue = PatternFill("solid", fgColor="1F4E78")
    white = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    normal = Font(name="Arial", size=10, color="1F1F1F")
    red_fill, red_font = PatternFill("solid", fgColor="FDE9E7"), Font(name="Arial", size=10, bold=True, color="C00000")
    pending_fill, pending_font = PatternFill("solid", fgColor="FFF2CC"), Font(name="Arial", size=10, bold=True, color="7F6000")
    gray, red = Side(style="thin", color="D9D9D9"), Side(style="thin", color="E26B6B")
    for cell in sheet[1]:
        cell.fill, cell.font = blue, white
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = normal
            cell.alignment = Alignment(vertical="center", wrap_text=cell.column >= 5)
            cell.border = Border(left=gray, right=gray, top=gray, bottom=gray)
    for cell in sheet[1]: cell.border = Border(left=gray, right=gray, top=gray, bottom=gray)
    for row_number in wrong_rows:
        for cell in sheet[row_number]:
            cell.fill, cell.font = red_fill, red_font
            cell.border = Border(left=red, right=red, top=red, bottom=red)
    for row_number in range(2, sheet.max_row + 1):
        if sheet.cell(row_number, 9).value == "待确认":
            for cell in sheet[row_number]: cell.fill, cell.font = pending_fill, pending_font
    for row in range(2, sheet.max_row + 1): sheet.cell(row, 2).number_format = "@"
    for index, width in enumerate([7, 16, 11, 15, 28, 12, 18, 28, 12, 30], 1): sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes, sheet.auto_filter.ref, sheet.sheet_view.showGridLines = "A2", f"A1:J{sheet.max_row}", False
    sheet.row_dimensions[1].height = 30
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main() -> None:
    args = parse_args()
    codes = load_codes(args.codes)
    if not codes: raise ValueError("款号文件中未识别到款号")
    plan = read_plan(args.plan, codes)
    actual_data = json.loads(args.actual.read_text(encoding="utf-8"))
    rows, wrong_rows = build_rows(codes, plan, actual_data)
    write_workbook(rows, wrong_rows, args.output)
    pending = sum(row[8] == "待确认" for row in rows)
    print(json.dumps({"styles": len(codes), "checks": len(rows), "correct": len(rows) - len(wrong_rows) - pending,
                      "wrong": len(wrong_rows), "pending": pending, "output": str(args.output)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
